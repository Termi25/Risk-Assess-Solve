"""Import Google-Forms questionnaire responses (.xlsx) into ``StudentCase`` objects.

The teacher collects answers with a Google Form ("Formular pentru Identificarea
Riscului de Abandon Școlar"); Google exports the responses as an ``.xlsx`` where
each row is one student and each column is one question (in order, every header
prefixed by its question number — ``"7.a …"``, ``"8. b …"``, ``"16. a …"``).

This module reads such a workbook and does two jobs:

1. **Column mapping** — resolves each header onto the app's questionnaire schema
   by its *question number*, so small wording / whitespace differences between
   forms don't break the import. Helper columns some teachers add by hand
   (``"… -cod"``, ``"Scor total"``, ``"Clasificare risc total"``) are ignored.

2. **Value normalisation** — Google-Form option labels are free text and rarely
   match the app's category vocabulary exactly (``"Familie monoparentală"`` vs
   ``"Monoparental"``, ``"Studii gimnaziale"`` vs ``"Gimnazial"``), Likert
   answers arrive as bare integers (``4`` → ``"Pozitivă"``), and the module
   average is a qualifier (``"Suficient"`` → ``6.0``). Every model-driving field
   is mapped onto the *exact* string the scoring engine expects, so an imported
   case scores and reports identically to one typed into the GUI. Without this,
   the encoder silently maps an unknown label to code 0 and mis-scores the case.

The module deliberately imports neither Qt nor the ML stack, so it stays cheap to
import and can be unit-tested in isolation.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .models import StudentCase


class ExcelImportError(ValueError):
    """User-facing (Romanian) error raised when a workbook cannot be imported."""


# --- header -> questionnaire key resolution ---------------------------------
# Keyed off the question number rather than the (variable) wording. A sub-question
# marker is a single trailing letter ("7.a", "8. b", "16. a").
_QNUM_TO_KEY: dict[str, str] = {
    "1": "full_name",
    "2": "birth_date",
    "3": "student_class",
    "4": "school_name",
    "5": "sex",
    "6": "residential_environment",
    "7": "family_situation",
    "7a": "family_situation_other",
    "8a": "mother_education",
    "8b": "father_education",
    "9": "unexcused_absences_3m",
    "10": "excused_absences_3m",
    "11": "extracurricular_participation",
    "12": "previous_module_average",
    "13": "low_grades_details",
    "14": "school_attitude",
    "15": "disciplinary_sanctions",
    "16": "school_feeling",
    "16a": "school_feeling_other",
    "17": "school_support_goal",
}

# Number + optional single-letter sub-marker at the very start of a header.
_QNUM_RE = re.compile(r"^\s*(\d+)\s*\.?\s*(?:([a-z])(?![a-z]))?")


def _strip_diacritics(text: str) -> str:
    nfkd = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in nfkd if not unicodedata.combining(ch))


def _norm(value: object) -> str:
    """Lower-cased, diacritic-free, whitespace-collapsed text for matching."""
    text = _strip_diacritics(str(value if value is not None else "")).lower()
    return re.sub(r"\s+", " ", text).strip()


def _header_key(header: object) -> str | None:
    """Questionnaire key for a column header, or ``None`` if the column is skipped."""
    norm = _norm(header)
    if not norm:
        return None
    # Manual helper columns added by some teachers: "… -cod", "Scor total",
    # "Clasificare risc total".
    if re.search(r"\bcod\b", norm) or "scor total" in norm or "clasificare" in norm:
        return None
    if "marcaj de timp" in norm or "timestamp" in norm:
        return "timestamp"
    match = _QNUM_RE.match(norm)
    if not match:
        return None
    qnum = match.group(1) + (match.group(2) or "")
    return _QNUM_TO_KEY.get(qnum)


# --- scalar coercions -------------------------------------------------------
def _to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_int(value: object, default: int = 0) -> int:
    try:
        text = str(value).replace(",", ".").strip()
        if not text:
            return default
        return int(round(float(text)))
    except (TypeError, ValueError):
        return default


def _to_float(value: object) -> float | None:
    try:
        text = str(value).replace(",", ".").strip()
        if not text:
            return None
        return float(text)
    except (TypeError, ValueError):
        return None


_TRIVIAL = {"", "-", "nu", "n/a", "na", "nu e cazul", "nu este cazul", "niciuna", "nimic"}


def _is_trivial(text: str) -> bool:
    return _norm(text) in _TRIVIAL


# --- value normalisers (onto the model's exact category vocabulary) ---------
def _norm_sex(value: object) -> str:
    n = _norm(value)
    if "femin" in n:
        return "Feminin"
    if "mascul" in n:
        return "Masculin"
    return "Altul / prefer să nu spun"


def _norm_environment(value: object) -> str:
    return "Rural" if "rural" in _norm(value) else "Urban"


def _norm_family(value: object) -> str:
    n = _norm(value)
    if "monoparent" in n or "divort" in n or "separ" in n:
        # Divorced/separated households are treated as single-parent for the
        # model (the closest category it was trained on).
        return "Monoparental"
    if "tutore" in n or "plasament" in n or "maternal" in n or "bunic" in n:
        return "Tutore / plasament"
    if "ambii" in n or "impreuna" in n or "casatorit" in n:
        return "Ambii părinți"
    if not n:
        return "Ambii părinți"
    return "Altă situație"


def _norm_education(value: object) -> str:
    n = _norm(value)
    if not n or "nu se aplica" in n or "necunoscut" in n or "nu stiu" in n:
        return "Necunoscut"
    if "fara studii" in n or "fara scoala" in n:
        return "Primar"
    if "postlice" in n:
        return "Postliceal"
    if "lice" in n:
        return "Liceal"
    if "gimnaz" in n:
        return "Gimnazial"
    if "primar" in n:
        return "Primar"
    if any(t in n for t in ("univers", "superioar", "facultate", "licenta", "master", "doctor")):
        return "Universitar"
    return "Necunoscut"


def _norm_participation(value: object) -> str:
    n = _norm(value)
    if not n or "nu particip" in n or "deloc" in n or n == "nu":
        return "Nu"
    if "ocazional" in n or "uneori" in n or "rar" in n or "cateodata" in n:
        return "Ocazional"
    return "Da, frecvent"


def _norm_average(value: object) -> float:
    number = _to_float(value)
    if number is not None and 1.0 <= number <= 10.0:
        return round(number, 1)
    n = _norm(value)
    if "foarte bine" in n or n == "fb":
        return 9.5
    if "excelent" in n:
        return 10.0
    if "insuficient" in n or n == "i":
        return 4.0
    if "suficient" in n or n == "s":
        return 6.0
    if "bine" in n or n == "b":
        return 8.0
    return 7.0


def _likert(value: object) -> float | None:
    """A bare linear-scale answer as a number (higher = more favourable)."""
    number = _to_float(value)
    return number if number is not None else None


def _norm_attitude(value: object) -> str:
    scale = _likert(value)
    if scale is not None:
        if scale < 2.5:
            return "Negativă"
        if scale < 3.5:
            return "Neutră"
        return "Pozitivă"
    n = _norm(value)
    if "pozit" in n:
        return "Pozitivă"
    if "negativ" in n:
        return "Negativă"
    return "Neutră"


def _norm_sanctions(value: object) -> str:
    n = _norm(value)
    if not n or n == "nu" or "niciun" in n or "nicio" in n:
        return "Nu"
    if "avertis" in n:
        return "Avertismente"
    if any(t in n for t in ("sanctiun", "exmatricul", "eliminare", "mutare disciplinar")):
        return "Sancțiuni"
    # A bare "Da" confirms a disciplinary event of unknown severity — treat it as
    # the milder "Avertismente" rather than overstating it as a formal sanction.
    return "Avertismente"


def _norm_feeling(value: object) -> str:
    n = _norm(value)
    if not n:
        return "Neutru"
    if any(t in n for t in ("stres", "anxi", "coples", "presiune", "panica", "frica", "speriat")):
        return "Stresat"
    if any(t in n for t in ("izol", "singur", "exclus", "respins", "marginaliz")):
        return "Izolat"
    if any(t in n for t in ("bine", "motivat", "implicat", "fericit", "bucur", "placut", "confortabil", "multumit")):
        return "Bine"
    if "neutru" in n or "normal" in n or "asa si asa" in n:
        return "Neutru"
    return "Altul"


def _norm_support(value: object) -> str:
    scale = _likert(value)
    if scale is not None:
        if scale < 2.5:
            return "Nu"
        if scale < 3.5:
            return "Parțial"
        return "Da"
    n = _norm(value)
    if "partial" in n or "uneori" in n or "putin" in n:
        return "Parțial"
    if n == "nu" or "deloc" in n:
        return "Nu"
    if "da" in n:
        return "Da"
    return "Parțial"


# Which normaliser handles each questionnaire key. Free-text keys (name, dates,
# details) are copied verbatim.
_NORMALISERS = {
    "sex": _norm_sex,
    "residential_environment": _norm_environment,
    "family_situation": _norm_family,
    "mother_education": _norm_education,
    "father_education": _norm_education,
    "extracurricular_participation": _norm_participation,
    "previous_module_average": _norm_average,
    "school_attitude": _norm_attitude,
    "disciplinary_sanctions": _norm_sanctions,
    "school_feeling": _norm_feeling,
    "school_support_goal": _norm_support,
}
_INT_KEYS = {"unexcused_absences_3m", "excused_absences_3m"}
_TEXT_KEYS = {
    "full_name", "student_class", "school_name",
    "family_situation_other", "low_grades_details", "school_feeling_other",
}


def _split_full_name(full_name: str) -> tuple[str, str]:
    """Split "Surname Given ..." into (surname, given names). Mirrors the GUI."""
    parts = [p for p in full_name.split() if p]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def _compose_observation(raw: dict[str, object]) -> str:
    """Free-text observation fed to the NLP engine.

    Includes the original (un-normalised) 'how do you feel at school' answer so
    descriptive replies ("Stresat / anxios") reach the emotional-stress signal.
    """
    parts: list[str] = []
    for key, label in (
        ("family_situation_other", "Situație familială - detalii"),
        ("low_grades_details", "Note sub 5 - detalii"),
        ("school_feeling", "Cum se simte la școală"),
        ("school_feeling_other", "Cum se simte la școală - detalii"),
        ("additional_notes", "Observații suplimentare"),
    ):
        text = _to_text(raw.get(key))
        if text and not _is_trivial(text):
            parts.append(f"{label}: {text}")
    return "\n".join(parts).strip()


def _row_to_case(raw: dict[str, object]) -> StudentCase:
    """Turn one raw questionnaire row (keyed by questionnaire key) into a case."""
    answers: dict[str, object] = {}
    timestamp = raw.get("timestamp")
    answers["timestamp"] = (
        timestamp.isoformat(sep=" ", timespec="seconds")
        if isinstance(timestamp, datetime)
        else _to_text(timestamp)
    )
    for key, value in raw.items():
        if key in ("timestamp",):
            continue
        if key in _NORMALISERS:
            answers[key] = _NORMALISERS[key](value)
        elif key in _INT_KEYS:
            answers[key] = _to_int(value)
        elif key == "birth_date":
            answers[key] = _to_text(value)  # ISO date; parsed downstream for age
        else:  # free-text keys
            answers[key] = _to_text(value)

    observation = _compose_observation(raw)
    features = dict(answers)
    features["observation_text"] = observation

    surname, name = _split_full_name(str(answers.get("full_name", "")))
    return StudentCase(
        name=name,
        surname=surname,
        student_grade=str(answers.get("student_class", "")).strip(),
        school_name=str(answers.get("school_name", "")).strip(),
        is_urban=answers.get("residential_environment") == "Urban",
        observation_text=observation,
        features=features,
    )


def _pick_sheet(workbook):
    """Choose the questionnaire-responses sheet.

    Prefers a Google "responses" sheet by name, then the sheet whose header row
    maps the most questionnaire columns. This skips manual working sheets that
    carry extra "-cod"/score columns (and often real names) in favour of the
    canonical export.
    """
    candidates = []
    for sheet in workbook.worksheets:
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        keys = {k for k in (_header_key(h) for h in header) if k}
        if {"full_name", "unexcused_absences_3m"} <= keys:
            name = _norm(sheet.title)
            prefers = any(t in name for t in ("raspuns", "response", "formular", "form"))
            candidates.append((prefers, len(keys), sheet))
    if not candidates:
        raise ExcelImportError(
            "Fișierul nu conține o foaie cu răspunsuri de chestionar recognoscibilă "
            "(lipsesc coloanele numerotate ale întrebărilor)."
        )
    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return candidates[0][2]


def load_cases_from_excel(path: str | Path) -> list[StudentCase]:
    """Read a Google-Forms responses ``.xlsx`` and return one case per student.

    Raises ``ExcelImportError`` with a user-facing (Romanian) message on any
    problem (missing file, wrong type, unreadable workbook, no recognisable
    responses sheet). Rows without a name and without any absence data are
    skipped as empty.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise ExcelImportError(f"Fișierul nu există: {file_path}")
    if file_path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ExcelImportError(
            "Sunt acceptate doar fișiere Excel (.xlsx) exportate din Google Forms."
        )

    try:
        workbook = load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises a variety of types
        raise ExcelImportError(
            f"Fișierul nu a putut fi citit ca registru Excel: {exc}"
        ) from exc

    try:
        sheet = _pick_sheet(workbook)
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ExcelImportError("Foaia de răspunsuri este goală.")

        header = rows[0]
        col_to_key: dict[int, str] = {}
        for index, cell in enumerate(header):
            key = _header_key(cell)
            if key and key not in col_to_key.values():
                col_to_key[index] = key

        cases: list[StudentCase] = []
        for row in rows[1:]:
            raw = {key: row[index] for index, key in col_to_key.items() if index < len(row)}
            # Skip blank trailing rows: no name and no absence figures.
            if not _to_text(raw.get("full_name")) and raw.get("unexcused_absences_3m") in (None, ""):
                continue
            cases.append(_row_to_case(raw))
    finally:
        workbook.close()

    if not cases:
        raise ExcelImportError("Nu s-au găsit răspunsuri de elevi în fișier.")
    return cases
